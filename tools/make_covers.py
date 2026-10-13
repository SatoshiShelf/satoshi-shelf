#!/usr/bin/env python3
import json, os, zipfile
from PIL import Image, ImageDraw, ImageFont, ImageOps

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "assets", "covers")
os.makedirs(OUT, exist_ok=True)
import subprocess
P = json.loads(subprocess.check_output(["node", "-e", "global.window={};require(process.argv[1]);process.stdout.write(JSON.stringify(window.PRODUCTS))", os.path.join(ROOT, "assets", "products.js")]))

W, H = 800, 450
FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FM = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
def F(p, s): return ImageFont.truetype(p, s)

CB = (64, 92, 736, 392)  # content box


def lerp(a, b, t): return tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3))


def base():
    img = Image.new("RGB", (W, H))
    d = ImageDraw.Draw(img)
    top, bot = (26, 31, 42), (12, 15, 21)
    for y in range(H):
        d.line([(0, y), (W, y)], fill=lerp(top, bot, y / H))
    for x in range(0, W, 26):
        for y in range(0, H, 26):
            d.point((x, y), fill=(38, 44, 56))
    ov = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(ov)
    od.ellipse([-180, -250, 470, 270], fill=(247, 147, 26, 26))
    od.ellipse([540, 300, 940, 580], fill=(56, 189, 248, 14))
    img = Image.alpha_composite(img.convert("RGBA"), ov).convert("RGB")
    return img


def window(fname):
    img = base()
    sh = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    sd = ImageDraw.Draw(sh)
    sd.rounded_rectangle([58, 62, 762, 420], radius=16, fill=(0, 0, 0, 120))
    img = Image.alpha_composite(img.convert("RGBA"), sh).convert("RGB")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([48, 44, 752, 404], radius=14, fill=(18, 21, 28), outline=(40, 46, 58))
    d.rounded_rectangle([48, 44, 752, 80], radius=14, fill=(31, 35, 45))
    d.rectangle([48, 66, 752, 80], fill=(31, 35, 45))
    for i, c in enumerate([(244, 91, 91), (245, 196, 81), (74, 208, 122)]):
        cx = 70 + i * 20
        d.ellipse([cx - 5, 57, cx + 5, 67], fill=c)
    d.text((124, 55), fname, font=F(FR, 15), fill=(150, 161, 176))
    return img, d


def fit_text(d, text, font, maxw):
    if d.textlength(text, font=font) <= maxw:
        return text
    while text and d.textlength(text + "\u2026", font=font) > maxw:
        text = text[:-1]
    return text + "\u2026"


def cover_xlsx(path):
    img, d = window(os.path.basename(path))
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i >= 8: break
        rows.append([("" if c is None else str(c)) for c in r[:5]])
    x0, y0, x1, y1 = CB
    d.rounded_rectangle([x0 - 4, y0 - 4, x1 + 4, y1 + 4], radius=8, fill=(244, 246, 249))
    rh = 26
    numw = 34
    ncols = 5
    cw = (x1 - x0 - numw) / ncols
    fh, fc = F(FB, 14), F(FR, 14)
    for ri, row in enumerate(rows):
        yy = y0 + ri * rh
        d.rectangle([x0, yy, x0 + numw, yy + rh], fill=(226, 232, 240))
        d.text((x0 + 9, yy + 5), str(ri + 1), font=fc, fill=(100, 116, 139))
        for ci in range(ncols):
            cx = x0 + numw + ci * cw
            cell = row[ci] if ci < len(row) else ""
            if ri == 0:
                d.rectangle([cx, yy, cx + cw, yy + rh], fill=(247, 147, 26))
                d.text((cx + 8, yy + 5), fit_text(d, cell, fh, cw - 16), font=fh, fill=(32, 20, 0))
            else:
                if ri % 2 == 1:
                    d.rectangle([cx, yy, cx + cw, yy + rh], fill=(248, 250, 252))
                d.text((cx + 8, yy + 5), fit_text(d, cell, fc, cw - 16), font=fc, fill=(30, 36, 48))
            d.line([(cx, yy), (cx, yy + rh)], fill=(210, 217, 226))
        d.line([(x0, yy + rh), (x1, yy + rh)], fill=(226, 232, 240))
    return img


def cover_docx(path):
    img, d = window(os.path.basename(path))
    import docx
    doc = docx.Document(path)
    title = None
    paras = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if not t: continue
        nm = (para.style.name or "").lower() if para.style else ""
        if nm.startswith(("heading", "title")) and title is None:
            title = t
        else:
            paras.append(t)
        if len(paras) >= 7: break
    x0, y0, x1, y1 = CB
    d.rounded_rectangle([x0, y0 - 6, x1, y1 + 6], radius=8, fill=(250, 251, 253))
    d.rectangle([x0, y0 - 6, x0 + 6, y1 + 6], fill=(247, 147, 26))
    px = x0 + 34
    py = y0 + 12
    if title:
        t = fit_text(d, title, F(FB, 22), x1 - px - 24)
        d.text((px, py), t, font=F(FB, 22), fill=(24, 28, 38))
        py += 40
    f = F(FR, 14)
    for p in paras:
        words = p.split()
        line = ""
        for w in words:
            if d.textlength(line + " " + w, font=f) > x1 - px - 24:
                d.text((px, py), line, font=f, fill=(90, 100, 116)); py += 21; line = w
                if py > y1 - 20: break
            else:
                line = (line + " " + w).strip()
        if line and py <= y1 - 20:
            d.text((px, py), line, font=f, fill=(90, 100, 116)); py += 21
        if py > y1 - 20: break
    return img


def cover_pptx(path):
    img, d = window(os.path.basename(path))
    from pptx import Presentation
    prs = Presentation(path)
    title = ""; body = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame: continue
            t = sh.text_frame.text.strip()
            if not t: continue
            lines = [l.strip() for l in t.split("\n") if l.strip()]
            if not title:
                title = lines[0]; lines = lines[1:]
            body += lines
        break
    x0, y0, x1, y1 = CB
    d.rounded_rectangle([x0, y0, x1, y1], radius=10, fill=(14, 17, 24))
    d.rectangle([x0, y0, x1, y0 + 8], fill=(247, 147, 26))
    px = x0 + 34
    py = y0 + 40
    d.text((px, py), fit_text(d, title or "Untitled", F(FB, 27), x1 - px - 30), font=F(FB, 27), fill=(240, 243, 248))
    py += 56
    f = F(FR, 16)
    for b in body[:6]:
        d.ellipse([px, py + 6, px + 7, py + 13], fill=(247, 147, 26))
        d.text((px + 20, py), fit_text(d, b, f, x1 - px - 40), font=f, fill=(190, 198, 210))
        py += 32
    return img


def cover_zip(path):
    img, d = window(os.path.basename(path))
    import zipfile
    z = zipfile.ZipFile(path)
    names = [n for n in z.namelist() if not n.endswith("/") and "__MACOSX" not in n]
    py = [n for n in names if n.endswith(".py")]
    lines = []
    if py:
        try:
            with z.open(py[0]) as fh:
                lines = fh.read().decode("utf-8", "replace").splitlines()[:18]
        except Exception:
            lines = []
    x0, y0, x1, y1 = CB
    d.rounded_rectangle([x0, y0, x1, y1], radius=8, fill=(13, 16, 22))
    fm = F(FM, 14)
    fn = F(FM, 13)
    yy = y0 + 10
    for i, ln in enumerate(lines):
        d.text((x0 + 12, yy), str(i + 1).rjust(2), font=fn, fill=(80, 89, 104))
        d.text((x0 + 42, yy), fit_text(d, ln.rstrip(), fm, x1 - x0 - 56), font=fm, fill=(176, 196, 220))
        yy += 19
        if yy > y1 - 16: break
    return img


def cover_html(pid):
    img, d = window(pid + ".html")
    src = os.path.join(ROOT, "previews", "img", pid + ".png")
    shot = Image.open(src).convert("RGB")
    x0, y0, x1, y1 = CB
    cw, ch = x1 - x0, y1 - y0
    fit = ImageOps.fit(shot, (cw, ch), method=Image.LANCZOS, centering=(0.5, 0.0))
    img.paste(fit, (x0, y0))
    d.rectangle([x0, y0, x1, y1], outline=(40, 46, 58))
    return img


def make(p):
    ext = p["file"].rsplit(".", 1)[-1].lower()
    fp = os.path.join(ROOT, p["file"])
    try:
        if ext == "html": return cover_html(p["id"])
        if ext == "xlsx": return cover_xlsx(fp)
        if ext == "docx": return cover_docx(fp)
        if ext == "pptx": return cover_pptx(fp)
        if ext == "zip": return cover_zip(fp)
    except Exception as e:
        print("FAIL", p["id"], e)
        return None
    return None


n = 0
for p in P:
    im = make(p)
    if im is None:
        continue
    im.save(os.path.join(OUT, p["id"] + ".jpg"), "JPEG", quality=84, optimize=True)
    n += 1
print("covers written:", n, "of", len(P))
